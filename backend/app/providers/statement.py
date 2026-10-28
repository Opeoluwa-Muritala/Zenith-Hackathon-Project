"""Bounded statement parsers; uploaded files are never persisted."""

import csv
import io
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation

from app.providers.base import ProviderError, ProviderTransaction

MAX_UPLOAD_BYTES = 2_000_000
MAX_ROWS = 10_000


class StatementUploadProvider:
    """Parse common CSV, XLSX and text-based PDF statement layouts."""

    def parse(self, data: bytes, filename: str = "statement.csv") -> list[ProviderTransaction]:
        if len(data) > MAX_UPLOAD_BYTES:
            raise ProviderError("Statement exceeds 2 MB")
        suffix = filename.lower().rsplit(".", 1)[-1]
        if suffix == "xlsx":
            return self._parse_xlsx(data)
        if suffix == "pdf":
            return self._parse_pdf(data)
        return self._parse_csv(data)

    def _parse_csv(self, data: bytes) -> list[ProviderTransaction]:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ProviderError("Statement must be UTF-8 CSV") from exc
        reader = csv.DictReader(io.StringIO(text))
        required = {"date", "amount_minor", "direction", "narration"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ProviderError("Missing required CSV columns")
        out = []
        for index, row in enumerate(reader):
            if index >= MAX_ROWS:
                raise ProviderError("Statement exceeds row limit")
            direction = row["direction"].lower()
            if direction not in {"credit", "debit"}:
                raise ProviderError("Invalid direction")
            amount = int(row["amount_minor"])
            if amount <= 0:
                raise ProviderError("Amount must be positive")
            out.append(
                ProviderTransaction(
                    row.get("id") or str(index),
                    datetime.fromisoformat(row["date"].replace("Z", "+00:00")),
                    amount,
                    direction,
                    row["narration"][:500],
                    row.get("category") or None,
                    int(row["balance_after_minor"]) if row.get("balance_after_minor") else None,
                    "upload",
                    8000,
                )
            )
        return out

    def _parse_xlsx(self, data: bytes) -> list[ProviderTransaction]:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise ProviderError("Invalid XLSX statement") from exc
        if not rows:
            raise ProviderError("Missing statement rows")
        headers = [str(value or "").strip().lower() for value in rows[0]]
        records = []
        for values in rows[1:]:
            records.append(dict(zip(headers, values, strict=False)))
        return self._parse_records(records)

    def _parse_pdf(self, data: bytes) -> list[ProviderTransaction]:
        try:
            from pypdf import PdfReader

            text = "\n".join(
                page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages
            )
        except Exception as exc:
            raise ProviderError("Invalid PDF statement") from exc
        records = []
        for index, line in enumerate(text.splitlines()):
            parts = [part.strip() for part in line.split("|")]
            if len(parts) >= 4:
                records.append(
                    {
                        "date": parts[0],
                        "amount_minor": parts[1],
                        "direction": parts[2],
                        "narration": parts[3],
                        "id": str(index),
                    }
                )
        return self._parse_records(records)

    def _parse_records(self, records: list[dict]) -> list[ProviderTransaction]:
        out = []
        for index, row in enumerate(records[: MAX_ROWS + 1]):
            if index >= MAX_ROWS:
                raise ProviderError("Statement exceeds row limit")
            try:
                posted = row.get("date")
                posted_at = (
                    posted
                    if isinstance(posted, datetime)
                    else datetime.fromisoformat(str(posted).replace("Z", "+00:00"))
                )
                if posted_at.tzinfo is None:
                    posted_at = posted_at.replace(tzinfo=UTC)
                direction = str(row.get("direction") or "").lower()
                amount = row.get("amount_minor")
                if amount in (None, ""):
                    credit, debit = row.get("credit"), row.get("debit")
                    direction = "credit" if credit not in (None, "") else "debit"
                    amount = credit if direction == "credit" else debit
                if amount in (None, ""):
                    raise ValueError
                raw_amount = str(amount).replace(",", "")
                amount_minor = (
                    int((Decimal(raw_amount) * 100).to_integral())
                    if isinstance(amount, str) and "." in raw_amount
                    else int(raw_amount)
                )
                narration = str(
                    row.get("narration") or row.get("description") or row.get("details") or ""
                )
                if direction not in {"credit", "debit"} or amount_minor <= 0 or not narration:
                    raise ValueError
            except (InvalidOperation, TypeError, ValueError, OverflowError) as exc:
                raise ProviderError("Invalid statement row") from exc
            out.append(
                ProviderTransaction(
                    str(row.get("id") or index),
                    posted_at,
                    amount_minor,
                    direction,
                    narration[:500],
                    row.get("category") or None,
                    None,
                    "upload",
                    8000,
                )
            )
        if not out:
            raise ProviderError("Missing statement rows")
        return out
